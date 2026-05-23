"""
Attendance routes — manual entry, live feed, attendance log, export, and AI insights.
Includes the SocketIO handler for real-time face recognition.
"""
import os
import time
import base64
import json
from datetime import datetime
from io import BytesIO

import numpy as np
import cv2
import dlib
import face_recognition
import requests
import pytz
from flask import render_template, request, redirect, url_for, flash, session, send_file, jsonify, current_app
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
import gc

from app.attendance import attendance_bp
from app.extensions import (
    db_connect, get_db, get_distinct_from_db, socketio,
    known_face_encodings, known_face_roll_numbers,
    last_logged_time
)
import app.extensions as ext
from app.face_utils import eye_aspect_ratio, predictor
from app.auth.routes import login_required, role_required
import config


# --- Helper: Get attendance data for display (present/absent) ---

def get_attendance_data_for_display(filters):
    local_tz = pytz.timezone(config.TIMEZONE)
    conn = get_db()

    all_students_query = "SELECT roll_no, name, branch, semester FROM students WHERE 1=1"
    student_query_params = []

    if filters.get('branch') and filters['branch'] != 'all':
        all_students_query += " AND branch = ?"
        student_query_params.append(filters['branch'])

    if filters.get('semester') and filters['semester'] != 'all':
        all_students_query += " AND semester = ?"
        try:
            student_query_params.append(int(filters['semester']))
        except (ValueError, TypeError):
            pass

    all_eligible_students = conn.execute(all_students_query, student_query_params).fetchall()

    present_logs_query = "SELECT log_id, student_roll_no, student_name, period_name, prof_name, timestamp FROM attendance_logs WHERE 1=1"
    present_log_params = []

    if filters.get('period_id') and filters['period_id'] != 'all':
        period_record = conn.execute("SELECT period_name, prof_name FROM schedule WHERE period_id = ?", (filters['period_id'],)).fetchone()
        if period_record:
            present_logs_query += " AND period_name = ?"
            present_log_params.append(period_record['period_name'])
        else:
            present_log_params.append("NON_EXISTENT_PERIOD")

    if filters.get('date'):
        present_logs_query += " AND date(timestamp) = ?"
        present_log_params.append(filters['date'])

    present_logs_raw = conn.execute(present_logs_query, present_log_params).fetchall()

    present_students_map = {}
    for log in present_logs_raw:
        present_students_map[log['student_roll_no']] = log

    final_attendance_records = []

    period_prof_name_for_absent = 'N/A'
    if filters.get('period_id') and filters['period_id'] != 'all':
        period_record = conn.execute("SELECT prof_name FROM schedule WHERE period_id = ?", (filters['period_id'],)).fetchone()
        if period_record:
            period_prof_name_for_absent = period_record['prof_name']

    for student in all_eligible_students:
        roll_no = student['roll_no']

        if roll_no in present_students_map:
            log = present_students_map[roll_no]
            log_dict = dict(log)
            log_dict['status'] = 'Present'

            if log_dict['timestamp']:
                try:
                    dt_naive = datetime.strptime(str(log_dict['timestamp']), '%Y-%m-%d %H:%M:%S')
                    dt_aware_local = local_tz.localize(dt_naive)
                    log_dict['timestamp'] = dt_aware_local.strftime('%Y-%m-%d %H:%M:%S')
                except ValueError:
                    log_dict['timestamp'] = str(log_dict['timestamp'])
            else:
                log_dict['timestamp'] = 'N/A'

            log_dict['student_branch'] = student['branch'] if student['branch'] else 'N/A'
            log_dict['student_semester'] = student['semester'] if student['semester'] else 'N/A'

            final_attendance_records.append(log_dict)
        else:
            if (filters.get('period_id') and filters['period_id'] != 'all') and filters.get('date'):
                if filters.get('status') == 'absent' or filters.get('status') == 'all':
                    final_attendance_records.append({
                        'log_id': None,
                        'student_roll_no': student['roll_no'],
                        'student_name': student['name'],
                        'period_name': conn.execute("SELECT period_name FROM schedule WHERE period_id = ?", (filters['period_id'],)).fetchone()['period_name'] if filters['period_id'] != 'all' else 'N/A',
                        'prof_name': period_prof_name_for_absent,
                        'timestamp': 'ABSENT',
                        'status': 'Absent',
                        'student_branch': student['branch'] if student['branch'] else 'N/A',
                        'student_semester': student['semester'] if student['semester'] else 'N/A'
                    })

    if filters.get('status') and filters['status'] != 'all':
        final_attendance_records = [
            record for record in final_attendance_records if record['status'].lower() == filters['status'].lower()
        ]

    final_attendance_records.sort(key=lambda x: x['student_name'])
    return final_attendance_records


def get_filtered_attendance_data(filters):
    conn = get_db()
    log_entries = []

    base_query = "SELECT log_id, student_roll_no, student_name, period_name, prof_name, timestamp FROM attendance_logs WHERE 1=1"
    query_params = []

    if filters.get('branch'):
        base_query += " AND student_roll_no IN (SELECT roll_no FROM students WHERE branch = ?)"
        query_params.append(filters['branch'])
    if filters.get('semester'):
        base_query += " AND student_roll_no IN (SELECT roll_no FROM students WHERE semester = ?)"
        query_params.append(filters['semester'])
    if filters.get('period_name'):
        base_query += " AND period_name = ?"
        query_params.append(filters['period_name'])
    if filters.get('date'):
        base_query += " AND date(timestamp) = ?"
        query_params.append(filters['date'])

    base_query += " ORDER BY timestamp DESC"
    present_students_raw = conn.execute(base_query, query_params).fetchall()

    present_roll_nos = set()
    for entry in present_students_raw:
        log_dict = dict(entry)
        student_info = conn.execute("SELECT branch, semester FROM students WHERE roll_no = ?", (log_dict['student_roll_no'],)).fetchone()
        log_dict['student_branch'] = student_info['branch'] if student_info else 'N/A'
        log_dict['student_semester'] = student_info['semester'] if student_info else 'N/A'
        log_dict['status'] = 'Present'
        log_entries.append(log_dict)
        present_roll_nos.add(log_dict['student_roll_no'])

    if filters.get('period_name') and filters.get('date'):
        all_students_query = "SELECT roll_no, name, branch, semester FROM students WHERE 1=1"
        student_params = []
        if filters.get('branch'):
            all_students_query += " AND branch = ?"
            student_params.append(filters['branch'])
        if filters.get('semester'):
            all_students_query += " AND semester = ?"
            student_params.append(filters['semester'])

        all_eligible_students = conn.execute(all_students_query, student_params).fetchall()

        for student in all_eligible_students:
            if student['roll_no'] not in present_roll_nos:
                log_entries.append({
                    'student_roll_no': student['roll_no'],
                    'student_name': student['name'],
                    'period_name': filters['period_name'],
                    'prof_name': 'N/A',
                    'timestamp': filters['date'],
                    'student_branch': student['branch'],
                    'student_semester': student['semester'],
                    'status': 'Absent'
                })
    return log_entries


# --- Routes ---

@attendance_bp.route('/manual_attendance', methods=['GET', 'POST'])
@login_required
@role_required(roles=['admin', 'professor'])
def manual_attendance():
    local_tz = pytz.timezone(config.TIMEZONE)
    conn = db_connect()

    filter_branch = request.args.get('branch')
    filter_semester = request.args.get('semester')
    filter_period_id_from_url = request.args.get('period_id')

    period_id_for_marking = request.form.get('period_id')
    form_data_for_template = request.form if request.method == 'POST' else request.args

    if request.method == 'POST':
        action = request.form.get('action')

        if action == 'mark_attendance':
            selected_roll_nos = request.form.getlist('roll_nos')
            current_filter_branch = request.form.get('current_filter_branch')
            current_filter_semester = request.form.get('current_filter_semester')
            current_filter_period_id = request.form.get('current_filter_period_id')

            if not selected_roll_nos:
                flash("Please select at least one student to mark as present.", "warning")
                conn.close()
                return redirect(url_for('attendance.manual_attendance',
                                        branch=current_filter_branch,
                                        semester=current_filter_semester,
                                        period_id=current_filter_period_id))

            if not period_id_for_marking or period_id_for_marking == 'all':
                flash("Please select a specific period.", "warning")
                conn.close()
                return redirect(url_for('attendance.manual_attendance',
                                        branch=current_filter_branch,
                                        semester=current_filter_semester,
                                        period_id=current_filter_period_id))

            period = conn.execute("SELECT period_name, prof_name FROM schedule WHERE period_id = ?", (period_id_for_marking,)).fetchone()
            if not period:
                flash("Selected period not found.", "danger")
                conn.close()
                return redirect(url_for('attendance.manual_attendance',
                                        branch=current_filter_branch,
                                        semester=current_filter_semester,
                                        period_id=current_filter_period_id))

            period_name, prof_name = period['period_name'], period['prof_name']
            logged_count = 0

            for roll_no in selected_roll_nos:
                student = conn.execute("SELECT name FROM students WHERE roll_no = ?", (roll_no,)).fetchone()
                if student:
                    student_name = student['name']
                    local_now = datetime.now(local_tz)
                    timestamp_str = local_now.strftime("%Y-%m-%d %H:%M:%S")
                    today_str = local_now.strftime("%Y-%m-%d")

                    existing_log = conn.execute(
                        "SELECT 1 FROM attendance_logs WHERE student_roll_no = ? AND period_name = ? AND date(timestamp) = ?",
                        (roll_no, period_name, today_str)
                    ).fetchone()

                    if not existing_log:
                        conn.execute(
                            "INSERT INTO attendance_logs (student_roll_no, student_name, period_name, prof_name, timestamp) VALUES (?, ?, ?, ?, ?)",
                            (roll_no, student_name, period_name, prof_name, timestamp_str)
                        )
                        logged_count += 1

            conn.commit()
            conn.close()

            if logged_count > 0:
                flash(f"Successfully marked attendance for {logged_count} student(s) in {period_name}.", "success")
            else:
                flash("No new attendance was logged (students may have already been marked present).", "info")

            return redirect(url_for('attendance.manual_attendance',
                                    branch=current_filter_branch,
                                    semester=current_filter_semester,
                                    period_id=current_filter_period_id))

    students_query = "SELECT roll_no, name, branch, semester FROM students WHERE 1=1"
    student_query_params = []

    if filter_branch and filter_branch != 'all':
        students_query += " AND branch = ?"
        student_query_params.append(filter_branch)

    if filter_semester and filter_semester != 'all':
        students_query += " AND semester = ?"
        try:
            student_query_params.append(int(filter_semester))
        except (ValueError, TypeError):
            filter_semester = None

    students_query += " ORDER BY name"
    students_list = conn.execute(students_query, student_query_params).fetchall()

    available_branches = get_distinct_from_db('branch', 'students')
    available_semesters = get_distinct_from_db('semester', 'students')
    schedule_list = conn.execute(
        "SELECT period_id, period_name, start_time, end_time, branch, semester FROM schedule ORDER BY start_time"
    ).fetchall()

    conn.close()

    return render_template(
        'attendance/manual_attendance.html',
        students=students_list,
        schedule=schedule_list,
        available_branches=available_branches,
        available_semesters=available_semesters,
        selected_branch=filter_branch,
        selected_semester=filter_semester,
        selected_period_id=filter_period_id_from_url,
        form_data=form_data_for_template
    )


@attendance_bp.route('/log', methods=['GET'])
@login_required
@role_required(roles=['admin', 'professor'])
def view_attendance_log():
    filter_branch = request.args.get('branch')
    filter_semester = request.args.get('semester')
    filter_period_id = request.args.get('period_id')
    filter_date_str = request.args.get('date')
    filter_status = request.args.get('status', 'all')

    current_filters = {
        'branch': filter_branch,
        'semester': filter_semester,
        'period_id': filter_period_id,
        'date': filter_date_str,
        'status': filter_status
    }

    log_entries = get_attendance_data_for_display(current_filters)

    conn = get_db()
    available_branches = get_distinct_from_db('branch', 'students')
    available_semesters = get_distinct_from_db('semester', 'students')
    schedule_periods = conn.execute(
        "SELECT period_id, period_name, start_time, end_time FROM schedule ORDER BY period_name"
    ).fetchall()

    return render_template(
        'attendance/log.html',
        logs=log_entries,
        schedule_periods=schedule_periods,
        available_branches=available_branches,
        available_semesters=available_semesters,
        filters=current_filters
    )


@attendance_bp.route('/delete_log_entry/<int:log_id>', methods=['POST'])
@login_required
@role_required(roles=['admin', 'professor'])
def delete_log_entry(log_id):
    conn = db_connect()
    try:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM attendance_logs WHERE log_id = ?", (log_id,))
        conn.commit()
        flash(f"Attendance log entry {log_id} deleted successfully!", "success")
    except Exception as e:
        flash(f"Error deleting log entry {log_id}: {e}", "danger")
    finally:
        conn.close()
    return redirect(url_for('attendance.view_attendance_log'))


@attendance_bp.route('/live')
@login_required
@role_required(roles=['admin', 'professor'])
def live_attendance():
    return render_template('attendance/live.html')


# --- Gemini API Routes ---

@attendance_bp.route('/generate_period_description', methods=['POST'])
@login_required
def generate_period_description():
    data = request.get_json()
    period_name = data.get('period_name')
    professor_name = data.get('professor_name')

    if not period_name or not professor_name:
        return jsonify({'status': 'error', 'message': 'Period name and professor name are required.'}), 400

    prompt = f"Generate a brief, engaging description (approx. 2-3 sentences) for a university course named '{period_name}' taught by Professor '{professor_name}'."

    try:
        api_key = os.environ.get("GEMINI_API_KEY")
        if not api_key:
            return jsonify({'status': 'error', 'message': 'GEMINI_API_KEY is not set.'}), 500

        api_url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={api_key}"
        payload = {
            "contents": [{"role": "user", "parts": [{"text": prompt}]}],
            "generationConfig": {"temperature": 0.7, "maxOutputTokens": 100}
        }

        response = requests.post(api_url, headers={'Content-Type': 'application/json'}, json=payload)
        response.raise_for_status()
        result = response.json()

        if result and result.get('candidates') and result['candidates'][0].get('content') and result['candidates'][0]['content'].get('parts'):
            generated_text = result['candidates'][0]['content']['parts'][0]['text']
            return jsonify({'status': 'success', 'description': generated_text})
        else:
            return jsonify({'status': 'error', 'message': 'Failed to get a valid response from AI.'}), 500

    except requests.exceptions.RequestException as e:
        return jsonify({'status': 'error', 'message': f'Error contacting AI: {e}'}), 500
    except Exception as e:
        return jsonify({'status': 'error', 'message': f'Something unexpected happened: {e}'}), 500


@attendance_bp.route('/generate_attendance_insights', methods=['POST'])
@login_required
def generate_attendance_insights():
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        return jsonify({'status': 'error', 'message': 'AI Service is not configured.'}), 500

    filters = request.get_json()
    log_entries = get_filtered_attendance_data(filters)

    if not log_entries:
        return jsonify({'status': 'error', 'message': 'No data found for the selected filters.'})

    present_count = sum(1 for log in log_entries if log['status'] == 'Present')
    absent_count = sum(1 for log in log_entries if log['status'] == 'Absent')
    total_students = present_count + absent_count

    if total_students == 0:
        return jsonify({'status': 'error', 'message': 'No students found for this filter combination.'})

    present_percentage = (present_count / total_students) * 100
    absent_students_list = [f"{log['student_name']} ({log['student_roll_no']})" for log in log_entries if log['status'] == 'Absent']

    prompt = f"""As an academic assistant, analyze the following daily attendance report:
    - Course/Period: {filters.get('period_name', 'N/A')}
    - Date: {filters.get('date', 'N/A')}
    - Total Students: {total_students}, Present: {present_count}, Absent: {absent_count}
    - Absent Students: {', '.join(absent_students_list) if absent_students_list else 'None'}
    - Attendance Percentage: {present_percentage:.2f}%

    Provide a brief summary with actionable suggestions. Use Markdown format."""

    api_url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-pro:generateContent?key={api_key}"
    payload = {"contents": [{"parts": [{"text": prompt}]}]}

    try:
        response = requests.post(api_url, json=payload)
        response.raise_for_status()
        generated_text = response.json()['candidates'][0]['content']['parts'][0]['text']
        return jsonify({'status': 'success', 'insights': generated_text.strip()})
    except Exception as e:
        return jsonify({'status': 'error', 'message': f'Error from AI service: {e}'}), 500


# --- Export ---

@attendance_bp.route('/export_attendance', methods=['GET'])
@login_required
@role_required(roles=['admin', 'professor'])
def export_attendance():
    filter_branch = request.args.get('branch')
    filter_semester = request.args.get('semester')
    filter_period_id = request.args.get('period_id')
    filter_date_str = request.args.get('date')
    filter_status = request.args.get('status', 'all')

    current_filters = {
        'branch': filter_branch,
        'semester': filter_semester,
        'period_id': filter_period_id,
        'date': filter_date_str,
        'status': filter_status
    }

    log_entries_for_excel = get_attendance_data_for_display(current_filters)

    if not log_entries_for_excel:
        flash("No attendance records found for the selected filters to export.", "warning")
        return redirect(url_for('attendance.view_attendance_log', **current_filters))

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance_Report"

    headers_map = {
        "ROLL NO": 'student_roll_no',
        "STUDENT NAME": 'student_name',
        "BRANCH": 'student_branch',
        "SEMESTER": 'student_semester',
        "PERIOD": 'period_name',
        "PROFESSOR": 'prof_name',
        "STATUS": 'status',
        "TIMESTAMP": 'timestamp'
    }

    header_titles = list(headers_map.keys())
    ws.append(header_titles)

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="F0F8FF", end_color="F0F8FF", fill_type="solid")
    for col_idx, cell in enumerate(ws[1]):
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.fill = header_fill
        ws.column_dimensions[cell.column_letter].width = 15

    for row_data in log_entries_for_excel:
        row_values = []
        for header_title in header_titles:
            key = headers_map[header_title]
            value = row_data.get(key, 'N/A')
            row_values.append(value)
        ws.append(row_values)

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                cell_value_str = str(cell.value) if cell.value is not None else ""
                if len(cell_value_str) > max_length:
                    max_length = len(cell_value_str)
            except:
                pass
        adjusted_width = (max_length + 2) * 1.1
        ws.column_dimensions[column].width = adjusted_width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = "attendance_report"
    if current_filters['branch'] and current_filters['branch'] != 'all':
        filename += f"_Branch_{current_filters['branch']}"
    if current_filters['semester'] and current_filters['semester'] != 'all':
        filename += f"_Sem_{current_filters['semester']}"
    if current_filters['date']:
        filename += f"_Date_{current_filters['date'].replace('-', '')}"
    if current_filters['status'] and current_filters['status'] != 'all':
        filename += f"_{current_filters['status']}"
    filename += f"_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )


session_frame_cache = {}

# --- SocketIO Handler ---

def register_socketio_handlers(socketio_instance):
    """Register SocketIO event handlers. Called from app factory."""

    @socketio_instance.on('image_from_client')
    def handle_image_from_client(data):
        start_total_time = time.time()
        detected_age_for_display = 'N/A'
        detected_gender_for_display = 'N/A'
        session_config = session.get('attendance_config')

        # Frame Skipping Cache logic per socket session
        sid = request.sid
        if sid not in session_frame_cache:
            session_frame_cache[sid] = {'count': 0, 'results': [], 'period': 'N/A'}
        
        cache = session_frame_cache[sid]
        cache['count'] += 1
        
        # Process 1 out of 3 frames, return cache for the others if cache is populated
        if cache['count'] % 3 != 0 and cache['results']:
            socketio_instance.emit('recognition_result', {
                'status': 'success',
                'results': cache['results'],
                'period': cache['period']
            })
            return

        try:
            img_data = base64.b64decode(data.split(',')[1])
            np_arr = np.frombuffer(img_data, np.uint8)
            img = cv2.imdecode(np_arr, cv2.IMREAD_COLOR)
            if img is None:
                raise ValueError("Could not decode image")
            
            # Dynamic Downscaling (target processing width = 320px) to optimize CPU and RAM usage
            h, w = img.shape[:2]
            target_w = 320
            if w > target_w:
                scale_factor = target_w / w
            else:
                scale_factor = 1.0
                
            small_img = cv2.resize(img, (0, 0), fx=scale_factor, fy=scale_factor)
            rgb_small_img = cv2.cvtColor(small_img, cv2.COLOR_BGR2RGB)

            face_locations = face_recognition.face_locations(rgb_small_img, model="hog")
            encodings = face_recognition.face_encodings(rgb_small_img, face_locations)

            conn = db_connect()
            try:
                # Query schedule periods ONCE per frame and reuse it
                periods = conn.execute("SELECT period_name, start_time, end_time, prof_name FROM schedule").fetchall()
                
                results = []
                for i, (encoding, location) in enumerate(zip(encodings, face_locations)):
                    name, roll_no, student_gender_from_db, student_age_from_db = None, None, 'N/A', 'N/A'
                    current_status = 'unknown'
                    liveness_status = 'N/A'

                    # Scale coordinates back to original size for frontend drawing and dlib predictor
                    (top, right, bottom, left) = location
                    orig_top = int(top / scale_factor)
                    orig_right = int(right / scale_factor)
                    orig_bottom = int(bottom / scale_factor)
                    orig_left = int(left / scale_factor)
                    orig_location = (orig_top, orig_right, orig_bottom, orig_left)

                    # --- Face Comparison ---
                    best_match_distance = 1.0
                    best_match_index = -1
                    if len(ext.known_face_encodings) > 0:
                        face_distances = face_recognition.face_distance(ext.known_face_encodings, encoding)
                        best_match_index = np.argmin(face_distances)
                        best_match_distance = face_distances[best_match_index]

                    # Resolve face identifier for per-face liveness tracking
                    is_recognized = best_match_distance < config.CONFIDENCE_THRESHOLD
                    if is_recognized:
                        roll_no = ext.known_face_roll_numbers[best_match_index]
                        student = conn.execute(
                            "SELECT name, branch, semester, gender, age FROM students WHERE roll_no = ?",
                            (roll_no,)
                        ).fetchone()

                        if student:
                            name = student['name']
                            student_gender_from_db = student['gender'] if student['gender'] else 'N/A'
                            student_age_from_db = student['age'] if student['age'] else 'N/A'
                            detected_gender_for_display = student_gender_from_db
                            detected_age_for_display = student_age_from_db
                        else:
                            name = "Unknown"
                            is_recognized = False
                    else:
                        name = "Unknown"

                    face_id = roll_no if is_recognized else f"unknown_{i}"

                    # --- Session / Schedule check ---
                    current_time_sec = time.time()
                    student_meets_criteria = False
                    current_period = None
                    current_prof_name = None
                    existing_log = None

                    if is_recognized and student and session_config:
                        config_branches = session_config.get('branches', [])
                        config_semester = session_config.get('semester')
                        config_subject = session_config.get('subject')

                        student_meets_criteria = (
                            (not config_branches or student['branch'] in config_branches) and
                            (config_semester is None or student['semester'] == config_semester) and
                            (not config_subject or student.get('subject') == config_subject)
                        )

                        if student_meets_criteria:
                            local_tz = pytz.timezone(config.TIMEZONE)
                            local_now_time = datetime.now(local_tz).time()
                            for period_rec in periods:
                                period_start = datetime.strptime(period_rec['start_time'], '%H:%M').time()
                                period_end = datetime.strptime(period_rec['end_time'], '%H:%M').time()
                                if period_start <= local_now_time <= period_end:
                                    current_period = period_rec['period_name']
                                    current_prof_name = period_rec['prof_name']
                                    break

                            if current_period:
                                today_str = datetime.now(local_tz).strftime("%Y-%m-%d")
                                existing_log = conn.execute(
                                    "SELECT 1 FROM attendance_logs WHERE student_roll_no = ? AND period_name = ? AND date(timestamp) = ?",
                                    (roll_no, current_period, today_str)
                                ).fetchone()

                    # --- Liveness Check Logic (Optimized, Per-Face State) ---
                    # If student is already logged, we can skip liveness to save CPU
                    if existing_log:
                        liveness_status = 'N/A'
                        current_status = 'already_logged'
                    else:
                        last_liveness_time = ext.last_liveness_check_times.get(face_id, 0)
                        in_cooldown = (current_time_sec - last_liveness_time) <= config.LIVENESS_CHECK_COOLDOWN

                        if predictor and not in_cooldown:
                            # Run landmarks on full resolution for precision (using orig_location rect)
                            gray_img = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
                            dlib_rect = dlib.rectangle(orig_left, orig_top, orig_right, orig_bottom)
                            shape = predictor(gray_img, dlib_rect)
                            shape = np.array([(shape.part(j).x, shape.part(j).y) for j in range(68)])

                            left_eye = shape[36:42]
                            right_eye = shape[42:48]
                            left_ear = eye_aspect_ratio(left_eye)
                            right_ear = eye_aspect_ratio(right_eye)
                            ear = (left_ear + right_ear) / 2.0

                            if ear < config.EYE_AR_THRESH:
                                ext.blink_counters[face_id] = ext.blink_counters.get(face_id, 0) + 1
                            else:
                                if ext.blink_counters.get(face_id, 0) >= config.EYE_AR_CONSEC_FRAMES:
                                    liveness_status = 'Blink Detected!'
                                    ext.last_liveness_check_times[face_id] = current_time_sec
                                ext.blink_counters[face_id] = 0

                            if liveness_status == 'N/A':
                                if ear > config.EYE_AR_THRESH * 1.5:
                                    liveness_status = 'Eyes Open'
                                elif ear < config.EYE_AR_THRESH / 2:
                                    liveness_status = 'Eyes Closed/Low EAR'
                                else:
                                    liveness_status = 'No Blink Yet'
                        else:
                            liveness_status = 'Liveness Cooldown' if in_cooldown else 'N/A'

                        # Determine overall status for recognized students
                        if is_recognized:
                            if not student:
                                current_status = 'db_mismatch'
                            elif not session_config:
                                current_status = 'session_not_configured'
                            elif not student_meets_criteria:
                                current_status = 'filter_mismatch'
                            elif not current_period:
                                current_status = 'no_active_period'
                            else:
                                # Candidate for logging
                                if roll_no in ext.last_logged_time and (current_time_sec - ext.last_logged_time[roll_no]) <= config.COOLDOWN_PERIOD:
                                    current_status = 'cooldown'
                                else:
                                    if liveness_status == 'Blink Detected!':
                                        local_tz = pytz.timezone(config.TIMEZONE)
                                        local_now_datetime = datetime.now(local_tz)
                                        timestamp_str = local_now_datetime.strftime("%Y-%m-%d %H:%M:%S")

                                        conn.execute(
                                            "INSERT INTO attendance_logs (student_roll_no, student_name, period_name, prof_name, timestamp) VALUES (?, ?, ?, ?, ?)",
                                            (roll_no, name, current_period, current_prof_name, timestamp_str)
                                        )
                                        conn.commit()
                                        ext.last_logged_time[roll_no] = current_time_sec
                                        current_status = 'logged'

                                        socketio_instance.emit('new_log_entry', {
                                            'name': name,
                                            'roll_no': roll_no,
                                            'gender': detected_gender_for_display,
                                            'age': detected_age_for_display,
                                            'time': local_now_datetime.strftime('%H:%M:%S')
                                        })
                                    else:
                                        # Liveness check still pending or cooldown from a previous session
                                        current_status = 'Liveness Check Failed'
                        else:
                            current_status = 'db_mismatch'

                    results.append({
                        'name': name,
                        'roll_no': roll_no,
                        'location': orig_location,
                        'status': current_status,
                        'gender': detected_gender_for_display,
                        'age': detected_age_for_display,
                        'distance': best_match_distance,
                        'liveness_status': liveness_status
                    })

                # Determine active period description if session_config is set
                active_period_name = 'N/A'
                if session_config:
                    local_tz = pytz.timezone(config.TIMEZONE)
                    local_now_time = datetime.now(local_tz).time()
                    for period_rec in periods:
                        period_start = datetime.strptime(period_rec['start_time'], '%H:%M').time()
                        period_end = datetime.strptime(period_rec['end_time'], '%H:%M').time()
                        if period_start <= local_now_time <= period_end:
                            active_period_name = period_rec['period_name']
                            break

                # Cache results for frame skipping
                cache['results'] = results
                cache['period'] = active_period_name

                socketio_instance.emit('recognition_result', {
                    'status': 'success',
                    'results': results,
                    'period': active_period_name
                })

            except Exception as e:
                print(f"Error in handle_image_from_client inner: {e}")
                import traceback
                traceback.print_exc()
                socketio_instance.emit('recognition_result', {'status': 'error', 'message': str(e), 'results': []})
            finally:
                conn.close()

                # Cleanup references to free memory immediately
                if 'img' in locals(): del img
                if 'small_img' in locals(): del small_img
                if 'rgb_small_img' in locals(): del rgb_small_img
                if 'gray_img' in locals(): del gray_img
                if 'np_arr' in locals(): del np_arr
                if 'img_data' in locals(): del img_data
                if 'encodings' in locals(): del encodings
                gc.collect()

            end_total_time = time.time()
            if current_app.debug:
                print(f"Total frame processing time: {(end_total_time - start_total_time)*1000:.2f} ms")

        except Exception as e:
            print(f"Top-level error in handle_image_from_client: {e}")
            import traceback
            traceback.print_exc()
            socketio_instance.emit('recognition_result', {'status': 'error', 'message': f"Top-level processing error: {e}", 'results': []})
